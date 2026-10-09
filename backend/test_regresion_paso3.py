"""
Pruebas de regresión - Paso 3: el Excel se leía con huecos silenciosos.

Fallas cubiertas (todas reproducidas antes de corregirlas):
  - Solo se leía la primera hoja; los reportes reales traen "Primera Opción" y "Segunda Opción".
  - Una fila vacía a mitad de la tabla cortaba la lectura.
  - IDs con puntos de miles ("1.110.487.315") se descartaban; los de más de 12 dígitos se truncaban.
  - Filas sin nombre, sin ID o con ID inválido desaparecían sin avisar.
  - Un .xlsx renombrado a .xls (o un archivo dañado) fallaba con un error interno.
  - Un Excel sin participantes se aceptaba y se guardaba un resultado vacío.

Datos inventados. Uso (desde la carpeta backend):
    python test_regresion_paso3.py
Termina con código 1 si alguna prueba falla. También es compatible con pytest.
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl

from app.document_grouper import DocumentRecord
from app.excel_reader import (
    MAX_ADVERTENCIAS, _split_identificacion, parse_excel_report, parse_excel_report_detailed,
)
from app.field_parser import ParsedFields
from app.matcher import compare
from app.report_generator import generar_reporte_pdf

ENCABEZADO = ["Identificación", "Nombre", "Estado"]


def libro(hojas: dict, encabezado=ENCABEZADO) -> bytes:
    """Libro .xlsx con la estructura del reporte real: filas de título y luego la tabla."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append([None, "Reporte de Inscripciones", None])
        ws.append([])
        ws.append(["Código Ficha", "1234567", None])
        ws.append(["Programa de Formación", "PROGRAMA DE PRUEBA", None])
        ws.append([])
        ws.append(encabezado)
        for f in filas:
            ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def ids(participantes):
    return [p.numero_documento for p in participantes]


# --------------------------------------------------------------------------- #
# Hojas
# --------------------------------------------------------------------------- #

def test_lee_todas_las_hojas_y_registra_de_cual_viene_cada_participante():
    data = libro({
        "Inscritos Primera Opción": [["CC - 1110487315", "ANA PEREZ", "Preinscrito"]],
        "Inscritos Segunda Opción": [["CC - 52123456", "LUIS GOMEZ", "Preinscrito"],
                                     ["TI - 1006521735", "JOSE DIAZ", "Preinscrito"]],
    })
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert ids(r.participantes) == ["1110487315", "52123456", "1006521735"]
    assert [p.hoja for p in r.participantes] == ["Inscritos Primera Opción"] + ["Inscritos Segunda Opción"] * 2
    assert r.advertencias == []


def test_segunda_hoja_vacia_es_normal_y_no_genera_advertencias():
    # Es lo que traen los reportes reales: la segunda hoja solo tiene el encabezado
    data = libro({"Inscritos Primera Opción": [["CC - 1110487315", "ANA PEREZ", "Preinscrito"]],
                  "Inscritos Segunda Opción": []})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert len(r.participantes) == 1 and r.advertencias == []


def test_hoja_sin_tabla_se_ignora_con_advertencia():
    wb = openpyxl.load_workbook(io.BytesIO(libro({"Datos": [["CC - 1110487315", "ANA PEREZ", "A"]]})))
    ws = wb.create_sheet("Notas")
    ws.append(["Texto libre sin tabla"])
    b = io.BytesIO(); wb.save(b)
    r = parse_excel_report_detailed(b.getvalue(), "x.xlsx")
    assert len(r.participantes) == 1
    assert any("Notas" in a and "ignorada" in a for a in r.advertencias), r.advertencias


# --------------------------------------------------------------------------- #
# Filas e identificaciones
# --------------------------------------------------------------------------- #

def test_fila_vacia_a_mitad_de_la_tabla_no_corta_la_lectura():
    data = libro({"H": [["CC - 1110487315", "ANA PEREZ", "A"], [None, None, None], ["CC - 52123456", "LUIS GOMEZ", "A"]]})
    assert ids(parse_excel_report(data, "x.xlsx")) == ["1110487315", "52123456"]


def test_formatos_de_identificacion():
    casos = {
        "CC - 1110487315": ("CC", "1110487315"),
        "CC-1.110.487.315": ("CC", "1110487315"),
        "1.110.487.315": (None, "1110487315"),
        "cc 52123456": ("CC", "52123456"),
        "TI - 1006521735": ("TI", "1006521735"),
        "1110487315.0": (None, "1110487315"),            # celda numérica
        "1.110487315E+09": (None, "1110487315"),         # notación científica
        "CC - 0012345678": ("CC", "0012345678"),         # ceros iniciales se conservan
    }
    for raw, esperado in casos.items():
        assert _split_identificacion(raw) == esperado, f"{raw!r} -> {_split_identificacion(raw)!r}"
    for invalido in ("CC - 1234", "SIN DOCUMENTO", "", None):
        assert _split_identificacion(invalido) == (None, None), invalido


def test_id_con_puntos_de_miles_ya_no_se_descarta():
    data = libro({"H": [["CC - 1.110.487.315", "ANA PEREZ", "A"]]})
    assert ids(parse_excel_report(data, "x.xlsx")) == ["1110487315"]


def test_id_de_mas_de_12_digitos_no_se_trunca_y_avisa():
    data = libro({"H": [["PEP - 123456789012345", "ANA PEREZ", "A"]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert ids(r.participantes) == ["123456789012345"]          # completo, no "123456789012"
    assert any("15 dígitos" in a and "no encontrado" in a for a in r.advertencias), r.advertencias


def test_fila_con_id_invalido_se_omite_pero_avisa():
    data = libro({"H": [["CC - 1234", "ANA PEREZ", "A"], ["CC - 52123456", "LUIS GOMEZ", "A"]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert ids(r.participantes) == ["52123456"]
    assert len(r.advertencias) == 1
    assert "fila 7" in r.advertencias[0] and "ANA PEREZ" in r.advertencias[0] and "NO aparece" in r.advertencias[0]


def test_fila_con_nombre_y_sin_id_avisa():
    data = libro({"H": [[None, "ANA PEREZ", "A"], ["CC - 52123456", "LUIS GOMEZ", "A"]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert ids(r.participantes) == ["52123456"]
    assert any("ANA PEREZ" in a and "no tiene identificación" in a for a in r.advertencias)


def test_fila_con_id_y_sin_nombre_se_conserva_y_avisa():
    data = libro({"H": [["CC - 1110487315", None, "A"]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert ids(r.participantes) == ["1110487315"] and r.participantes[0].nombre_reportado == ""
    assert any("no tiene nombre" in a for a in r.advertencias)


def test_documento_repetido_avisa():
    data = libro({"H": [["CC - 1110487315", "ANA PEREZ", "A"], ["CC - 1110487315", "ANA PEREZ", "A"]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert len(r.participantes) == 2
    assert any("repetido" in a and "fila 7" in a for a in r.advertencias), r.advertencias


def test_fila_de_totales_al_pie_se_ignora_sin_ruido():
    data = libro({"H": [["CC - 1110487315", "ANA PEREZ", "A"], ["Total", "1", None]]})
    r = parse_excel_report_detailed(data, "x.xlsx")
    assert len(r.participantes) == 1 and r.advertencias == []


def test_columnas_nombres_y_apellidos_separadas_se_unen():
    data = libro({"H": [["CC - 1110487315", "ANA MARIA", "PEREZ GOMEZ", "A"]]},
                 encabezado=["Identificación", "Nombres", "Apellidos", "Estado"])
    p = parse_excel_report(data, "x.xlsx")[0]
    assert p.nombre_reportado == "ANA MARIA PEREZ GOMEZ", p.nombre_reportado
    assert p.estado == "A"


def test_exceso_de_advertencias_se_resume():
    filas = [[f"CC - {i}", f"PERSONA {i}", "A"] for i in range(1000, 1040)] + [["CC - 52123456", "LUIS GOMEZ", "A"]]
    r = parse_excel_report_detailed(libro({"H": filas}), "x.xlsx")
    assert len(r.advertencias) == MAX_ADVERTENCIAS + 1
    assert r.advertencias[-1].startswith("... y 15 advertencia(s) más"), r.advertencias[-1]


def test_archivo_limpio_no_genera_advertencias():
    data = libro({"H": [["CC - 1110487315", "ANA PEREZ", "A"], ["TI - 1006521735", "JOSE DIAZ", "A"]]})
    assert parse_excel_report_detailed(data, "x.xlsx").advertencias == []


# --------------------------------------------------------------------------- #
# Archivos inválidos: siempre ValueError con mensaje claro (nunca un error interno)
# --------------------------------------------------------------------------- #

def test_xlsx_con_extension_xls_se_lee_igual():
    data = libro({"H": [["CC - 1110487315", "ANA PEREZ", "A"]]})
    assert ids(parse_excel_report(data, "reporte.xls")) == ["1110487315"]


def test_archivo_danado_da_valueerror_claro():
    for nombre in ("x.xls", "x.xlsx"):
        try:
            parse_excel_report(b"no soy un excel", nombre)
        except ValueError as exc:
            assert "No se pudo leer" in str(exc) and nombre in str(exc), str(exc)
        else:
            raise AssertionError(f"{nombre}: debia lanzar ValueError")


def test_excel_sin_tabla_o_sin_participantes_se_rechaza():
    sin_tabla = libro({"H": []}, encabezado=["Otra", "Cosa", "Distinta"])
    try:
        parse_excel_report(sin_tabla, "x.xlsx")
    except ValueError as exc:
        assert "encabezado" in str(exc)
    else:
        raise AssertionError("debia rechazar un Excel sin la tabla esperada")

    sin_filas = libro({"H": [["CC - 12", "ANA", "A"]]})       # la única fila es inválida
    try:
        parse_excel_report(sin_filas, "x.xlsx")
    except ValueError as exc:
        assert "ninguna fila con una identificación válida" in str(exc)
    else:
        raise AssertionError("debia rechazar un Excel sin participantes")


# --------------------------------------------------------------------------- #
# Cruce e informe
# --------------------------------------------------------------------------- #

def test_participante_sin_nombre_con_documento_encontrado_no_queda_como_coincide():
    p = parse_excel_report(libro({"H": [["CC - 1000000007", None, "A"]]}), "x.xlsx")
    doc = DocumentRecord(
        numero_documento="1000000007",
        fields=ParsedFields(numero_documento="1000000007", apellidos="GARCIA LOPEZ", nombres="MARIA ELENA",
                            fecha_nacimiento="1990-01-01", fuente="frente_clasico"),
        source_pages=[1],
    )
    fila = compare([doc], p)["resultados"][0]
    assert fila["estado"] == "nombre_no_coincide"
    assert "no trae nombre" in fila["observaciones"]


def test_informe_pdf_incluye_advertencias_y_tolera_caracteres_especiales():
    resultado = {
        "resumen": {"total_excel": 1, "coinciden": 1},
        "resultados": [], "sobrantes_en_pdf": [], "documentos_sin_numero_detectado": [],
        "advertencias_excel": ["Hoja 'A&B', fila 7: la identificación '<b>x' no es válida"],
    }
    pdf = generar_reporte_pdf("a.pdf", "b.xls", resultado, "admin")
    import pymupdf
    texto = "\n".join(pg.get_text() for pg in pymupdf.open(stream=pdf, filetype="pdf"))
    assert "Advertencias de lectura del Excel" in texto
    assert "A&B" in texto and "<b>x" in texto            # el texto del Excel se muestra tal cual, sin interpretarse


# --------------------------------------------------------------------------- #
# Ejecución como script
# --------------------------------------------------------------------------- #

if __name__ == "__main__":
    pruebas = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    fallos = 0
    for nombre, fn in pruebas:
        try:
            fn()
            print(f"[PASS] {nombre}")
        except Exception as exc:  # noqa: BLE001 - cualquier error cuenta como fallo de la prueba
            fallos += 1
            print(f"[FAIL] {nombre}\n       {type(exc).__name__}: {exc}")
    print(f"\n{len(pruebas) - fallos}/{len(pruebas)} pruebas correctas")
    sys.exit(1 if fallos else 0)
